from __future__ import annotations

import argparse
import json
from itertools import combinations, permutations
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
DEFAULT_GEN = HERE / "outputs" / "qwen2.5_3b_L24" / "generations" / "mental_health_800tok_diag.jsonl"
DEFAULT_OUT = HERE / "outputs" / "qwen2.5_3b_L24" / "pairs_plan.xlsx"

DEFAULT_CELLS = "baseline,global_ridge@0.2,global_ridge@0.3,local_manifold@0.2,local_manifold@0.3"
DEFAULT_TIERS = {"Collapse", "Striving", "Conflict"}

Cell = Tuple[str, Optional[float]]


def parse_cell(s: str) -> Cell:
    s = s.strip()
    if "@" in s:
        c, a = s.split("@", 1)
        a = a.strip()
        return (c.strip(), None if a.lower() in ("none", "base", "") else float(a))
    return (s, None)


def cell_label(cell: Cell) -> str:
    c, a = cell
    return c if a is None else f"{c}@{a:g}"


def _sidx(r: Dict) -> int:
    for k in ("sample_idx", "sample", "sample_id"):
        if k in r:
            try:
                return int(r[k])
            except (TypeError, ValueError):
                pass
    return 0


def _resp(r: Dict) -> str:
    for k in ("response", "generation", "completion", "text", "output"):
        v = r.get(k)
        if isinstance(v, str) and v.strip():
            return v
    return ""


def load_groups(path: Path, tiers: set, samples: set) -> Dict[Tuple[str, int], Dict[Cell, Dict]]:
    groups: Dict[Tuple[str, int], Dict[Cell, Dict]] = {}
    with open(path) as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            r = json.loads(line)
            tier = (r.get("prompt_metadata") or {}).get("tier")
            if tiers and tier not in tiers:
                continue
            s = _sidx(r)
            if samples and s not in samples:
                continue
            key = (r["prompt_id"], s)
            groups.setdefault(key, {})[(r.get("condition"), r.get("alpha"))] = r
    return groups


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--in", dest="infile", type=str, default=str(DEFAULT_GEN))
    p.add_argument("--out", dest="outfile", type=str, default=str(DEFAULT_OUT))
    p.add_argument("--cells", type=str, default=DEFAULT_CELLS,
                   help="Comma-separated cells (condition or condition@alpha).")
    p.add_argument("--tiers", type=str, default=",".join(sorted(DEFAULT_TIERS)),
                   help="Comma-separated tiers, or 'all'.")
    p.add_argument("--samples", type=str, default="0,1",
                   help="Comma-separated sample indices.")
    p.add_argument("--both-orders", action="store_true",
                   help="Emit each pair in BOTH A/B orders (760 rows) to cancel and "
                        "measure position bias.  Default: one row per pair, side randomized.")
    p.add_argument("--seed", type=int, default=42,
                   help="Seed for the random left/right assignment (reproducible).")
    p.add_argument("--max-resp-chars", type=int, default=4000,
                   help="Truncate response text in the sheet to this many chars.")
    args = p.parse_args()

    import random
    rng = random.Random(args.seed)

    cells = [parse_cell(c) for c in args.cells.split(",") if c.strip()]
    tiers = set() if args.tiers.strip().lower() == "all" else {t.strip() for t in args.tiers.split(",") if t.strip()}
    samples = {int(s) for s in args.samples.split(",") if s.strip()}

    infile = Path(args.infile)
    outfile = Path(args.outfile)
    outfile.parent.mkdir(parents=True, exist_ok=True)
    if not infile.exists():
        raise SystemExit(f"Input not found: {infile}")

    groups = load_groups(infile, tiers, samples)
    pair_list = list(combinations(cells, 2))

    rows = []
    pair_coverage = {(cell_label(a), cell_label(b)): 0 for a, b in pair_list}
    cid = 0
    for (pid, s), by_cell in sorted(groups.items()):
        md = next(iter(by_cell.values())).get("prompt_metadata") or {}
        tier = md.get("tier")
        cat = md.get("category")
        prompt_text = next(iter(by_cell.values())).get("prompt", "")
        for a, b in pair_list:
            rA0, rB0 = by_cell.get(a), by_cell.get(b)
            if rA0 is None or rB0 is None:
                continue
            pair_coverage[(cell_label(a), cell_label(b))] += 1
            if args.both_orders:
                orients = [(a, rA0, b, rB0), (b, rB0, a, rA0)]
            else:
                orients = [(a, rA0, b, rB0)] if rng.random() < 0.5 else [(b, rB0, a, rA0)]
            for (cellA, rowA, cellB, rowB) in orients:
                cid += 1
                rows.append({
                    "comparison_id": cid,
                    "prompt_id": pid,
                    "sample_idx": s,
                    "tier": tier,
                    "category": cat,
                    "prompt": prompt_text,
                    "response_A": _resp(rowA)[:args.max_resp_chars],
                    "response_B": _resp(rowB)[:args.max_resp_chars],
                    "cell_A": cell_label(cellA),
                    "cell_B": cell_label(cellB),
                })

    wb = Workbook()
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="404040")
    body_font = Font(name="Arial")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    def _header(ws, headers):
        for j, h in enumerate(headers, 1):
            c = ws.cell(1, j, h)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center")

    ws = wb.active
    ws.title = "to_judge"
    jh = ["comparison_id", "tier", "category", "prompt", "response_A", "response_B", "winner (A/B)"]
    _header(ws, jh)
    for i, row in enumerate(rows, 2):
        vals = [row["comparison_id"], row["tier"], row["category"], row["prompt"],
                row["response_A"], row["response_B"], ""]
        for j, (h, v) in enumerate(zip(jh, vals), 1):
            c = ws.cell(i, j, v)
            c.font = body_font
            c.alignment = wrap if h in ("prompt", "response_A", "response_B") else top
    for j, w in zip(range(1, len(jh) + 1), [12, 10, 14, 45, 60, 60, 12]):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("key")
    kh = ["comparison_id", "prompt_id", "sample_idx", "tier", "cell_A", "cell_B"]
    _header(ws2, kh)
    for i, row in enumerate(rows, 2):
        for j, h in enumerate(kh, 1):
            ws2.cell(i, j, row[h]).font = body_font
    for j, w in zip(range(1, len(kh) + 1), [12, 12, 10, 10, 20, 20]):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("pairs")
    _header(ws3, ["pair_index", "cell_A", "cell_B", "n_groups_covered"])
    for i, ((la, lb), n) in enumerate(pair_coverage.items(), start=2):
        for j, v in enumerate([i - 1, la, lb, n], 1):
            ws3.cell(i, j, v).font = body_font
    for j, w in enumerate([12, 20, 20, 18], 1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    ws3.freeze_panes = "A2"

    wb.save(outfile)
    print(f"groups: {len(groups)}")
    print(f"cells: {[cell_label(c) for c in cells]}")
    print(f"unordered pairs: {len(pair_list)}  ({'both orders' if args.both_orders else 'side randomized, seed=%d' % args.seed})")
    print(f"total rows to judge: {len(rows)}")
    print(f"wrote: {outfile}")


if __name__ == "__main__":
    main()
