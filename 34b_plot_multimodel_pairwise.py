"""Plot per-model round-robin win rates from the blind pairwise evaluation (Section 2.5.3).

Reads the win_rate_by_tier sheet of each model's judged workbook (18_llm_judge.py), pools the low tiers,
and writes pairwise_multimodel.{pdf,png}, LaTeX tables (overall and by tier), and a summary CSV.
Usage: python 34b_plot_multimodel_pairwise.py --in-dir all_mh_judged
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent

# (judged workbook filename, display name)
MODELS = [
    ("mh_eval_results_L4_qwen_3b_L24.xlsx",                "Qwen2.5-3B"),
    ("mh_llm_eval_results_qwen7b_L20.xlsx",                "Qwen2.5-7B"),
    ("mental_health_eval_qwen14b_L24_juedge_claude.xlsx",  "Qwen2.5-14B"),
    ("mh_llm_eval_results_llama8b_L16_claude.xlsx",        "Llama-3.1-8B"),
    ("mh_llm_eval_results_gemma_L29_claude.xlsx",          "Gemma-2-9B"),
    ("mh_llm_eval_results_gemma2_27b_judged_claude.xlsx",  "Gemma-2-27B"),
]
TIERS = ["Collapse", "Striving", "Conflict"]
ALPHAS = ["0.2", "0.3"]
CELLS = ["baseline"] + [f"{m}@{a}" for m in ("global_ridge", "local_manifold") for a in ALPHAS]
COLOR = {"baseline": "#7f7f7f", "global_ridge": "#d62728", "local_manifold": "#2ca02c"}
TEXTWIDTH_IN = 5.5


def load_model(path: Path):
    """Read {(tier, cell): (appearances, wins)} from win_rate_by_tier; exit if any cell is missing."""
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["win_rate_by_tier"].iter_rows(values_only=True))
    H = {h: i for i, h in enumerate(rows[0])}
    out = {}
    for r in rows[1:]:
        if r[H["tier"]] in TIERS and r[H["condition"]] in CELLS:
            out[(r[H["tier"]], r[H["condition"]])] = (int(r[H["appearances"]]), int(r[H["wins"]]))
    missing = [(t, c) for t in TIERS for c in CELLS if (t, c) not in out]
    if missing:
        raise SystemExit(f"{path.name}: missing {missing}")
    return out


def overall(counts, cell):
    """Win rate (%) pooled over all tiers, with total wins and appearances."""
    n = sum(counts[(t, cell)][0] for t in TIERS)
    w = sum(counts[(t, cell)][1] for t in TIERS)
    return 100.0 * w / n, w, n


def tier_rate(counts, tier, cell):
    n, w = counts[(tier, cell)]
    return 100.0 * w / n


def plot(data, out: Path):
    """Per model: dashed baseline line, ridge vs. manifold markers per alpha (filled 0.2, hollow 0.3)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    plt.rcParams.update({"font.size": 7, "xtick.labelsize": 7, "ytick.labelsize": 6,
                         "legend.fontsize": 6.5, "pdf.fonttype": 42, "ps.fonttype": 42})
    fig, ax = plt.subplots(figsize=(TEXTWIDTH_IN, 2.0))
    dx = {"0.2": -0.17, "0.3": 0.17}
    for i, (name, counts) in enumerate(data):
        ax.hlines(overall(counts, "baseline")[0], i - 0.4, i + 0.4,
                  color="0.45", ls="--", lw=0.9, zorder=2)
        for a in ALPHAS:
            x = i + dx[a]
            r = overall(counts, f"global_ridge@{a}")[0]
            m = overall(counts, f"local_manifold@{a}")[0]
            filled = a == "0.2"
            # Grey connector between the ridge and manifold rates at the same alpha.
            ax.plot([x, x], [r, m], color="0.78", lw=1.4, zorder=2, solid_capstyle="butt")
            ax.plot(x, r, marker="X", ms=5, ls="none", color=COLOR["global_ridge"], mew=1,
                    mfc=COLOR["global_ridge"] if filled else "white", zorder=3)
            ax.plot(x, m, marker="s", ms=4, ls="none", color=COLOR["local_manifold"], mew=1,
                    mfc=COLOR["local_manifold"] if filled else "white", zorder=3)
            ax.text(x, -2.5, f"α={a[1:]}", ha="center", va="top", fontsize=5.5, color="0.4")
    ax.set_xticks(range(len(data)))
    ax.set_xticklabels([n for n, _ in data])
    ax.tick_params(axis="x", length=0, pad=11)
    for x in [i + 0.5 for i in range(len(data) - 1)]:
        ax.axvline(x, color="0.9", lw=0.6, zorder=0)
    ax.set_xlim(-0.5, len(data) - 0.5)
    ax.set_ylim(0, 85)
    ax.set_ylabel("round-robin win rate (%)")
    ax.grid(axis="y", alpha=0.3, lw=0.4, zorder=0)
    handles = [
        Line2D([], [], color="0.45", ls="--", lw=0.9, label="baseline"),
        Line2D([], [], marker="X", ms=5, ls="none", color=COLOR["global_ridge"], label="global_ridge"),
        Line2D([], [], marker="s", ms=4, ls="none", color=COLOR["local_manifold"], label="local_manifold"),
        Line2D([], [], marker="o", ms=4, ls="none", color="0.3", label="filled: α=0.2"),
        Line2D([], [], marker="o", ms=4, ls="none", color="0.3", mfc="white", label="hollow: α=0.3"),
    ]
    ax.legend(handles=handles, loc="lower center", bbox_to_anchor=(0.5, 1.0), ncol=5,
              frameon=False, handletextpad=0.3, columnspacing=1.3)
    fig.tight_layout(pad=0.3)
    fig.savefig(out / "pairwise_multimodel.pdf", bbox_inches="tight")
    fig.savefig(out / "pairwise_multimodel.png", dpi=300, bbox_inches="tight")


def _tex_cell(c):
    """LaTeX column header for a cell."""
    if c == "baseline":
        return "baseline"
    m, a = c.split("@")
    return f"\\texttt{{{m.replace('_', chr(92) + '_')}}}@{a}"


def _row(vals, fmt):
    """Join values as LaTeX table cells, bolding the maximum."""
    best = max(vals)
    return " & ".join((f"\\textbf{{{fmt(v)}}}" if v == best else fmt(v)) for v in vals)


def write_tex(data, path: Path):
    """Write two LaTeX tables: overall win rate and win rate by tier."""
    hdr = " & ".join(_tex_cell(c) for c in CELLS)
    L = []
    L.append("% Generated by 34b_plot_multimodel_pairwise.py -- do not edit by hand")
    L.append("% Table 1: overall round-robin win rate (low tiers), wins/appearances in parentheses")
    L.append("\\begin{table}[t]\n\\centering\n\\small")
    L.append("\\caption{Blind pairwise evaluation on low-tier prompts: round-robin win rate (\\%) "
             "per condition, with wins/appearances. Rates are zero-sum within a model "
             "(the five conditions average 50\\%). Best per model in bold.}")
    L.append("\\label{tab:pairwise_overall}")
    L.append("\\resizebox{\\linewidth}{!}{%")
    L.append("\\begin{tabular}{l" + "c" * len(CELLS) + "}\n\\toprule")
    L.append(f"Model & {hdr} \\\\\n\\midrule")
    for name, counts in data:
        ov = [overall(counts, c) for c in CELLS]
        best = max(v[0] for v in ov)
        cells = []
        for r, w, n in ov:
            s = f"{r:.1f} ({w}/{n})"
            cells.append(f"\\textbf{{{s}}}" if r == best else s)
        L.append(f"{name} & " + " & ".join(cells) + " \\\\")
    L.append("\\bottomrule\n\\end{tabular}}\n\\end{table}\n")

    L.append("% Table 2: win rate by tier")
    L.append("\\begin{table}[t]\n\\centering\n\\small")
    L.append("\\caption{Pairwise win rate (\\%) by prompt tier. Best per row in bold.}")
    L.append("\\label{tab:pairwise_by_tier}")
    L.append("\\resizebox{\\linewidth}{!}{%")
    L.append("\\begin{tabular}{ll" + "c" * len(CELLS) + "}\n\\toprule")
    L.append(f"Model & Tier & {hdr} \\\\\n\\midrule")
    for k, (name, counts) in enumerate(data):
        for j, t in enumerate(TIERS):
            vals = [tier_rate(counts, t, c) for c in CELLS]
            first = f"\\multirow{{{len(TIERS)}}}{{*}}{{{name}}}" if j == 0 else ""
            n = counts[(t, "baseline")][0]
            L.append(f"{first} & {t} ($n$={n}) & " + _row(vals, lambda v: f"{v:.1f}") + " \\\\")
        if k < len(data) - 1:
            L.append("\\midrule")
    L.append("\\bottomrule\n\\end{tabular}}\n\\end{table}")
    L.append("% needs \\usepackage{booktabs,multirow,graphicx}")
    path.write_text("\n".join(L) + "\n")


def write_csv(data, path: Path):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["model", "tier", "condition", "appearances", "wins", "win_rate_pct"])
        for name, counts in data:
            for c in CELLS:
                r, wi, n = overall(counts, c)
                w.writerow([name, "all_low_tiers", c, n, wi, f"{r:.1f}"])
            for t in TIERS:
                for c in CELLS:
                    n, wi = counts[(t, c)]
                    w.writerow([name, t, c, n, wi, f"{100.0 * wi / n:.1f}"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in-dir", default=str(HERE / "all_mh_judged"))
    ap.add_argument("--out-dir", default=str(HERE / "outputs" / "figures_multimodel"))
    args = ap.parse_args()

    data = [(name, load_model(Path(args.in_dir) / f)) for f, name in MODELS]
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    plot(data, out)
    write_tex(data, out / "pairwise_multimodel_tables.tex")
    write_csv(data, out / "pairwise_multimodel_summary.csv")

    # Count (model, alpha) settings where local_manifold beats global_ridge / baseline.
    wins_vs_ridge = sum(overall(c, f"local_manifold@{a}")[0] > overall(c, f"global_ridge@{a}")[0]
                        for _, c in data for a in ALPHAS)
    wins_vs_base = sum(overall(c, f"local_manifold@{a}")[0] > overall(c, "baseline")[0]
                       for _, c in data for a in ALPHAS)
    n = len(data) * len(ALPHAS)
    print(f"local_manifold > global_ridge (same alpha): {wins_vs_ridge}/{n}")
    print(f"local_manifold > baseline:                  {wins_vs_base}/{n}")
    print(f"wrote: {out / 'pairwise_multimodel.pdf'}, pairwise_multimodel_tables.tex, pairwise_multimodel_summary.csv")


if __name__ == "__main__":
    main()
